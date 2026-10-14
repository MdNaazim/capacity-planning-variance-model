"""Builds Capacity_Planning_Variance_Model.xlsx (SYNTHETIC data, live Excel formulas).
Run: python build_model.py   (then open the .xlsx in Excel)
"""
import numpy as np
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo

rng = np.random.default_rng(21)
REGIONS = ["North", "South", "East", "West", "Central", "North-East"]
FAC_PER_REGION, MONTHS = 7, 12
F = "Arial"
blue = Font(name=F, color="0000FF"); black = Font(name=F); bold = Font(name=F, bold=True)
hdr_font = Font(name=F, bold=True, color="FFFFFF"); hdr_fill = PatternFill("solid", fgColor="1F3864")
thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
# ------------------------------------------------------------------ README
rd = wb.active; rd.title = "README"
text = [
    ("Capacity Planning & Performance Variance Model", True),
    ("All figures are SYNTHETIC (simulated for portfolio purposes; seed 21).", False),
    ("", False),
    ("Sheets", True),
    ("Targets   - lookup table: facility -> region and target utilisation (blue = editable input)", False),
    ("Data      - 504 facility-month records (42 facilities x 12 months). Blue = inputs; black = formulas", False),
    ("Summary   - region and monthly variance analysis built with SUMIFS / AVERAGEIFS / COUNTIFS", False),
    ("", False),
    ("Formulas used", True),
    ("VLOOKUP (Region and Target Utilisation pulled from Targets), SUMIFS, COUNTIFS, AVERAGEIFS, IF", False),
    ("Conditional formatting: Variance % colour scale, status colours, data bars on utilisation", False),
    ("", False),
    ("Definitions", True),
    ("Variance = Actual Output - Planned Capacity;  Variance % = Variance / Planned;  Utilisation = Actual / Planned", False),
    ("Gap vs Target = Utilisation - Target Utilisation;  Status: Under if gap < -5 pts, Over if gap > +5 pts, else On Track", False),
    ("Status band (+/-5 pts) is in Targets!F2 and can be changed.", False),
    ("", False),
    ("Pivot Table (build in 30 seconds)", True),
    ("Data sheet > click any cell in the table > Insert > PivotTable > new sheet. Rows: Region; Columns: Status;", False),
    ("Values: Sum of Planned Capacity, Sum of Actual Output. (openpyxl cannot create pivot caches, so this step is manual.)", False),
]
for i, (t, b) in enumerate(text, 1):
    rd.cell(i, 1, t).font = Font(name=F, bold=b, size=14 if i == 1 else 11)
rd.column_dimensions["A"].width = 120

# ------------------------------------------------------------------ Targets
tg = wb.create_sheet("Targets")
for c, h in enumerate(["Facility ID", "Region", "Target Utilisation", "Base Capacity (units/month)"], 1):
    x = tg.cell(1, c, h); x.font = hdr_font; x.fill = hdr_fill; x.alignment = Alignment(horizontal="center", wrap_text=True)
facilities = []
r = 2
for region in REGIONS:
    for k in range(FAC_PER_REGION):
        fid = f"F{len(facilities) + 1:02d}"
        target = float(rng.choice([0.85, 0.88, 0.90, 0.92]))
        base = int(rng.integers(900, 3200) // 50 * 50)
        facilities.append((fid, region, target, base))
        for c, v in enumerate([fid, region, target, base], 1):
            cell = tg.cell(r, c, v); cell.font = blue; cell.border = box
        tg.cell(r, 3).number_format = "0%"; tg.cell(r, 4).number_format = "#,##0"
        r += 1
tg["F1"] = "Status band (+/- pts)"; tg["F1"].font = bold
tg["F2"] = 0.05; tg["F2"].font = blue; tg["F2"].number_format = "0%"; tg["F2"].fill = PatternFill("solid", fgColor="FFFF00")
tg["F3"] = "Assumption: utilisation within +/-5 pts of target = On Track"; tg["F3"].font = Font(name=F, italic=True)
for col, w in zip("ABCDEF", [12, 14, 20, 26, 22, 55]): tg.column_dimensions[col].width = w
tg.freeze_panes = "A2"
LAST_T = 1 + len(facilities)

# ------------------------------------------------------------------ Data
ds = wb.create_sheet("Data")
heads = ["Month", "Facility ID", "Region", "Planned Capacity", "Actual Output", "Variance", "Variance %",
         "Utilisation", "Target Util.", "Gap vs Target", "Status"]
for c, h in enumerate(heads, 1):
    x = ds.cell(1, c, h); x.font = hdr_font; x.fill = hdr_fill; x.alignment = Alignment(horizontal="center", wrap_text=True)
region_bias = {"North": 0.00, "South": -0.09, "East": -0.02, "West": 0.04, "Central": -0.04, "North-East": -0.07}
season = [0.95, 0.95, 1.0, 1.0, 1.02, 1.03, 1.0, 1.0, 1.05, 1.08, 1.10, 1.06]
row = 2
for m in range(MONTHS):
    for fid, region, target, base in facilities:
        planned = int(round(base * season[m] / 10) * 10)
        util = np.clip(target + region_bias[region] + rng.normal(0, 0.05), 0.55, 1.12)
        actual = int(round(planned * util))
        ds.cell(row, 1, date(2026, m + 1, 1)).number_format = "mmm-yy"
        ds.cell(row, 2, fid); ds.cell(row, 4, planned); ds.cell(row, 5, actual)
        for c in (1, 2, 4, 5): ds.cell(row, c).font = blue
        ds.cell(row, 3, f"=VLOOKUP(B{row},Targets!$A$2:$D${LAST_T},2,FALSE)")
        ds.cell(row, 6, f"=E{row}-D{row}")
        ds.cell(row, 7, f"=IF(D{row}=0,0,F{row}/D{row})")
        ds.cell(row, 8, f"=IF(D{row}=0,0,E{row}/D{row})")
        ds.cell(row, 9, f"=VLOOKUP(B{row},Targets!$A$2:$D${LAST_T},3,FALSE)")
        ds.cell(row, 10, f"=H{row}-I{row}")
        ds.cell(row, 11, f'=IF(J{row}<-Targets!$F$2,"Under",IF(J{row}>Targets!$F$2,"Over","On Track"))')
        for c in (3, 6, 7, 8, 9, 10, 11): ds.cell(row, c).font = black
        for c in (4, 5, 6): ds.cell(row, c).number_format = "#,##0;(#,##0);-"
        for c in (7, 8, 9): ds.cell(row, c).number_format = "0.0%"
        ds.cell(row, 10).number_format = '0.0%;[Red]-0.0%'
        row += 1
LAST = row - 1
tbl = Table(displayName="CapacityData", ref=f"A1:K{LAST}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
ds.add_table(tbl)
ds.freeze_panes = "A2"
for col, w in zip("ABCDEFGHIJK", [10, 12, 13, 17, 15, 12, 12, 12, 12, 14, 12]): ds.column_dimensions[col].width = w
ds.row_dimensions[1].height = 30
# conditional formatting
ds.conditional_formatting.add(f"G2:G{LAST}", ColorScaleRule(start_type="num", start_value=-0.25, start_color="F8696B",
                              mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="num", end_value=0.15, end_color="63BE7B"))
ds.conditional_formatting.add(f"K2:K{LAST}", CellIsRule(operator="equal", formula=['"Under"'], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(color="9C0006")))
ds.conditional_formatting.add(f"K2:K{LAST}", CellIsRule(operator="equal", formula=['"Over"'], fill=PatternFill("solid", bgColor="FFEB9C"), font=Font(color="9C5700")))
ds.conditional_formatting.add(f"K2:K{LAST}", CellIsRule(operator="equal", formula=['"On Track"'], fill=PatternFill("solid", bgColor="C6EFCE"), font=Font(color="006100")))
ds.conditional_formatting.add(f"H2:H{LAST}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1.2, color="5B9BD5"))

# ------------------------------------------------------------------ Summary
sm = wb.create_sheet("Summary")
sm["A1"] = "Capacity Planning & Performance Variance Summary (synthetic data)"; sm["A1"].font = Font(name=F, bold=True, size=14)
sm["A2"] = "All values are formulas on the Data sheet. Change inputs (blue) on Data / Targets and this sheet updates."; sm["A2"].font = Font(name=F, italic=True)
R = lambda col: f"Data!${col}$2:${col}${LAST}"
h = ["Region", "Planned Capacity", "Actual Output", "Variance", "Variance %", "Utilisation", "Avg Target Util.", "Gap vs Target", "Under-target Facility-Months", "% Under"]
for c, t in enumerate(h, 1):
    x = sm.cell(4, c, t); x.font = hdr_font; x.fill = hdr_fill; x.alignment = Alignment(horizontal="center", wrap_text=True)
for i, region in enumerate(REGIONS):
    r = 5 + i
    sm.cell(r, 1, region)
    sm.cell(r, 2, f"=SUMIFS({R('D')},{R('C')},$A{r})")
    sm.cell(r, 3, f"=SUMIFS({R('E')},{R('C')},$A{r})")
    sm.cell(r, 4, f"=C{r}-B{r}")
    sm.cell(r, 5, f"=IF(B{r}=0,0,D{r}/B{r})")
    sm.cell(r, 6, f"=IF(B{r}=0,0,C{r}/B{r})")
    sm.cell(r, 7, f"=AVERAGEIFS({R('I')},{R('C')},$A{r})")
    sm.cell(r, 8, f"=F{r}-G{r}")
    sm.cell(r, 9, f'=COUNTIFS({R("C")},$A{r},{R("K")},"Under")')
    sm.cell(r, 10, f'=IF(COUNTIFS({R("C")},$A{r})=0,0,I{r}/COUNTIFS({R("C")},$A{r}))')
tr = 5 + len(REGIONS)
sm.cell(tr, 1, "Total")
for c, col in ((2, "B"), (3, "C"), (4, "D"), (9, "I")):
    sm.cell(tr, c, f"=SUM({col}5:{col}{tr-1})")
sm.cell(tr, 5, f"=IF(B{tr}=0,0,D{tr}/B{tr})"); sm.cell(tr, 6, f"=IF(B{tr}=0,0,C{tr}/B{tr})")
sm.cell(tr, 7, f"=AVERAGE({R('I')})"); sm.cell(tr, 8, f"=F{tr}-G{tr}")
sm.cell(tr, 10, f"=I{tr}/COUNTA({R('B')})")
for r in range(5, tr + 1):
    for c in range(1, 11):
        x = sm.cell(r, c); x.font = bold if r == tr else black; x.border = box
    for c in (2, 3, 4, 9): sm.cell(r, c).number_format = "#,##0;(#,##0);-"
    for c in (5, 6, 7, 8, 10): sm.cell(r, c).number_format = "0.0%;[Red]-0.0%"
sm.conditional_formatting.add(f"E5:E{tr-1}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="max", end_color="63BE7B"))
sm.conditional_formatting.add(f"H5:H{tr-1}", CellIsRule(operator="lessThan", formula=["-Targets!$F$2"], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(color="9C0006")))
sm.conditional_formatting.add(f"F5:F{tr-1}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1.2, color="5B9BD5"))

# monthly table
mr = tr + 3
sm.cell(mr - 1, 1, "Monthly performance").font = bold
for c, t in enumerate(["Month", "Planned Capacity", "Actual Output", "Variance", "Variance %", "Utilisation"], 1):
    x = sm.cell(mr, c, t); x.font = hdr_font; x.fill = hdr_fill; x.alignment = Alignment(horizontal="center", wrap_text=True)
for m in range(MONTHS):
    r = mr + 1 + m
    sm.cell(r, 1, date(2026, m + 1, 1)).number_format = "mmm-yy"; sm.cell(r, 1).font = blue
    sm.cell(r, 2, f"=SUMIFS({R('D')},{R('A')},$A{r})"); sm.cell(r, 3, f"=SUMIFS({R('E')},{R('A')},$A{r})")
    sm.cell(r, 4, f"=C{r}-B{r}"); sm.cell(r, 5, f"=IF(B{r}=0,0,D{r}/B{r})"); sm.cell(r, 6, f"=IF(B{r}=0,0,C{r}/B{r})")
    for c in range(1, 7):
        x = sm.cell(r, c); x.border = box
        if c > 1: x.font = black
    for c in (2, 3, 4): sm.cell(r, c).number_format = "#,##0;(#,##0);-"
    for c in (5, 6): sm.cell(r, c).number_format = "0.0%;[Red]-0.0%"
mend = mr + MONTHS
sm.conditional_formatting.add(f"E{mr+1}:E{mend}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="max", end_color="63BE7B"))
for col, w in zip("ABCDEFGHIJ", [14, 17, 15, 14, 12, 12, 15, 13, 20, 10]): sm.column_dimensions[col].width = w
sm.row_dimensions[4].height = 42

# charts
ch = BarChart(); ch.type = "col"; ch.title = "Utilisation vs Target by Region"; ch.y_axis.number_format = "0%"
ch.add_data(Reference(sm, min_col=6, min_row=4, max_row=tr - 1), titles_from_data=True)
ch.add_data(Reference(sm, min_col=7, min_row=4, max_row=tr - 1), titles_from_data=True)
ch.set_categories(Reference(sm, min_col=1, min_row=5, max_row=tr - 1)); ch.height, ch.width = 8, 15
ch.y_axis.scaling.min = 0.6; ch.y_axis.delete = False; ch.x_axis.delete = False
sm.add_chart(ch, "L4")
lc = LineChart(); lc.title = "Monthly Variance %"; lc.y_axis.number_format = "0%"
lc.add_data(Reference(sm, min_col=5, min_row=mr, max_row=mend), titles_from_data=True)
lc.set_categories(Reference(sm, min_col=1, min_row=mr + 1, max_row=mend)); lc.height, lc.width = 8, 15
lc.y_axis.delete = False; lc.x_axis.delete = False; lc.x_axis.number_format = "mmm"
sm.add_chart(lc, "L22")

wb.move_sheet("Summary", offset=-2)   # order: README, Summary, Targets, Data
wb.save("Capacity_Planning_Variance_Model.xlsx")
print("saved; data rows:", LAST - 1)
